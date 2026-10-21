from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "output/docs/Casita_Magma_Founding_Vision.docx"

INK = "201713"
CHARCOAL = "342820"
MAGMA = "A64222"
EMBER = "D0653D"
SAND = "F4EEE6"
PARCHMENT = "FBF8F3"
MUTED = "786D63"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in kwargs.items():
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def set_table_widths(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_font(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def paragraph(doc, text="", style=None, before=None, after=None, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        set_font(p.add_run(text))
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_font(p.add_run(text), size=10.6)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_font(p.add_run(text), size={1: 16, 2: 12.5, 3: 11.5}[level], color=MAGMA, bold=True)
    return p


def add_lead(doc, text):
    p = doc.add_paragraph(style="Magma Lead")
    set_font(p.add_run(text), size=13, color=CHARCOAL, italic=True)
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [9120])
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, SAND)
    set_cell_margins(cell, top=170, bottom=170, start=240, end=240)
    set_cell_border(cell, top={"val": "single", "sz": "12", "color": MAGMA}, bottom={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label.upper() + "  ")
    set_font(r, size=9.5, color=MAGMA, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.7, color=INK)
    paragraph(doc, after=3)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [2220, 6900])
    for label, value in rows:
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell, top={"val": "single", "sz": "4", "color": "D9CEC2"}, bottom={"val": "single", "sz": "4", "color": "D9CEC2"}, start={"val": "nil"}, end={"val": "nil"})
        set_cell_shading(cells[0], SAND)
        p = cells[0].paragraphs[0]
        set_font(p.add_run(label), size=10, color=MAGMA, bold=True)
        p = cells[1].paragraphs[0]
        set_font(p.add_run(value), size=10.4)
    paragraph(doc, after=3)


def add_phase_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [1920, 3250, 3950])
    header = table.rows[0].cells
    prevent_row_split(table.rows[0])
    for cell, label in zip(header, ["PHASE", "FOCUS", "PROOF OF LIFE"]):
        set_cell_shading(cell, CHARCOAL)
        set_cell_margins(cell, top=100, bottom=100)
        set_cell_border(cell, top={"val": "single", "sz": "4", "color": CHARCOAL}, bottom={"val": "single", "sz": "4", "color": CHARCOAL}, start={"val": "single", "sz": "4", "color": CHARCOAL}, end={"val": "single", "sz": "4", "color": CHARCOAL})
        p = cell.paragraphs[0]
        set_font(p.add_run(label), size=9.3, color="FFFFFF", bold=True)
    rows = [
        ("1. Gather", "Small recurring events, radio, and a founding cohort.", "People return; collaborators meet; culture begins to cohere."),
        ("2. Connect", "Mobile-first event rooms, colectivo spaces, and project needs.", "Useful connections happen beyond the founding circle."),
        ("3. Hold", "Participation-based keys, trusted moderation, and a living archive.", "Membership feels earned through care, not purchased status."),
        ("4. House", "Pop-ups, a temporary clubhouse, then a permanent home.", "A real community already needs and sustains the space."),
    ]
    for i, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for cell, value in zip(cells, values):
            set_cell_margins(cell, top=120, bottom=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "FFFFFF" if i % 2 == 0 else PARCHMENT)
            set_cell_border(cell, top={"val": "single", "sz": "4", "color": "D9CEC2"}, bottom={"val": "single", "sz": "4", "color": "D9CEC2"}, start={"val": "single", "sz": "4", "color": "D9CEC2"}, end={"val": "single", "sz": "4", "color": "D9CEC2"})
            p = cell.paragraphs[0]
            set_font(p.add_run(value), size=9.6, color=INK, bold=(cell == cells[0]))
    paragraph(doc, after=3)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.24

for level, size, before, after in [(1, 16, 18, 8), (2, 12.5, 13, 6), (3, 11.5, 10, 4)]:
    s = styles[f"Heading {level}"]
    s.font.name = "Aptos Display"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(MAGMA)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

lead = styles.add_style("Magma Lead", WD_STYLE_TYPE.PARAGRAPH)
lead.paragraph_format.space_before = Pt(2)
lead.paragraph_format.space_after = Pt(15)
lead.paragraph_format.line_spacing = 1.25

bullet = styles["List Bullet"]
bullet.base_style = styles["Normal"]
bullet.paragraph_format.left_indent = Inches(0.31)
bullet.paragraph_format.first_line_indent = Inches(-0.16)
bullet.paragraph_format.space_after = Pt(3)
bullet.paragraph_format.line_spacing = 1.18

# Header/footer.
header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header_p.add_run("CASITA MAGMA  /  FOUNDING VISION")
set_font(r, size=8.5, color=MUTED, bold=True)
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer_p.add_run("Open-hearted. Carefully held. Built through participation.")
set_font(r, size=8.5, color=MUTED, italic=True)

# Cover.
p = paragraph(doc, before=88, after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
set_font(p.add_run("A FOUNDING VISION"), size=10, color=MAGMA, bold=True)
p = paragraph(doc, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
set_font(p.add_run("CASITA MAGMA"), size=33, color=INK, bold=True)
p = paragraph(doc, after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
set_font(p.add_run("The digital home for a cultural house in Mexico City."), size=15, color=CHARCOAL, italic=True)
p = paragraph(doc, after=28, align=WD_ALIGN_PARAGRAPH.CENTER)
set_font(p.add_run("Magma is the network.\nCasita Magma is the home.\nKeys are how people enter, contribute, and become part of it."), size=13, color=MAGMA, bold=True)
add_callout(doc, "Core idea", "A community-run cultural commons where artists, collectivos, and curious people meet, make work, and carry real-world energy into a shared digital space.")
p = paragraph(doc, before=70, align=WD_ALIGN_PARAGRAPH.CENTER)
set_font(p.add_run("CDMX  •  2026"), size=9.5, color=MUTED, bold=True)
doc.add_page_break()

# 1.
add_heading(doc, "1. What Casita Magma is", 1)
add_lead(doc, "A place of inclusion where people want to be—yet not a place built for everyone.")
paragraph(doc, "Casita Magma is a digital home for a community-run cultural house in Mexico City. It exists to make relationships, creative work, and real-world gatherings easier to discover, deepen, and remember. It does not attempt to cover the city. It follows the people shaping it.")
add_callout(doc, "The distinction", "Magma is inclusive in who can belong, but intentional about what it hosts. Anyone can be welcomed; not every behavior, project, or partnership is aligned.")
add_heading(doc, "The culture it protects", 2)
for item in [
    "Curiosity, generosity, experimentation, and a real point of view.",
    "Local connection over extractive tourism, status-chasing, and follower culture.",
    "Participation and care over money, notoriety, or proximity to the right people.",
    "Autonomy for collectivos, alongside infrastructure that connects them.",
]:
    add_bullet(doc, item)
add_heading(doc, "Brand architecture", 2)
add_key_value_table(doc, [
    ("Magma", "The cultural network: collectivos, artists, projects, radio, gatherings, and shared momentum."),
    ("Casita Magma", "The home: the digital space now and, eventually, a physical cultural house in CDMX."),
    ("Keys", "A symbolic access system earned through presence, contribution, and trusted invitation."),
    ("La Mano", "A member’s visible record of how they help hold up the house—not popularity or follower count."),
])

# 2.
add_heading(doc, "2. Two paths inward", 1)
add_lead(doc, "Belonging is never for sale. It is built through participation.")
paragraph(doc, "The physical world matters to Magma, but it cannot be the only doorway. Someone who lives outside CDMX, cannot afford a ticket, or cannot be physically present should still be able to become a real part of the community.")
add_heading(doc, "Presence path", 2)
paragraph(doc, "Attend a gathering, help produce it, volunteer, bring documentation, host a workshop, play a radio set, or otherwise show up for the community.")
add_heading(doc, "Contribution path", 2)
paragraph(doc, "Participate from anywhere: make or share work, join radio, host a digital salon, help curate, write field notes, translate, document, contribute to a project, or be vouched for by trusted members.")
add_callout(doc, "Design principle", "Physical artifacts are optional souvenirs and keys to a particular moment—not proof that someone is more worthy of belonging.")
add_heading(doc, "Keys to the house", 2)
paragraph(doc, "A printed card, artist-made sticker, zine, poster, matchbook, or single-use code can carry a key. Redeeming it connects an event to a richer digital layer: its room, archive, collaborators, radio recordings, and future invitations. In time, the key can become literal: a key to Casita Magma’s physical door.")

# 3.
add_heading(doc, "3. La Mano: contribution, not clout", 1)
add_lead(doc, "A profile should show the mark someone leaves on the community—not how many people follow them.")
paragraph(doc, "La Mano is a contribution profile. It shows the ways someone helps hold up the house: what they share, organize, document, make possible, and bring to others. It is not a public ranking or a popularity contest.")
add_key_value_table(doc, [
    ("Shared", "Posts, field notes, art, images, and useful resources."),
    ("In conversation", "Thoughtful replies, responses, and sustained discussion."),
    ("Gathered", "Events attended and documented."),
    ("Made possible", "Sessions hosted, care work, production, moderation, or behind-the-scenes help."),
    ("Collaborated", "Shared projects and colectivo work, with credit visible across profiles."),
])
doc.add_page_break()
add_heading(doc, "Guardrails", 2)
for item in [
    "Weight meaningful actions more than raw reactions; limit rewards for repetitive engagement.",
    "Recognize unseen labor—setup, care, translation, accessibility, documentation, and moderation.",
    "Use La Mano to unlock trust and context, not a leaderboard or a hierarchy of worth.",
    "Let members show selected roles or badges without exposing an exact total if they prefer privacy.",
]:
    add_bullet(doc, item)

# 4.
add_heading(doc, "4. Collectivos own their spaces; Magma connects the edges", 1)
add_lead(doc, "Magma is the bridge and hosting space for new waves of thought and connection around art in Mexico City.")
paragraph(doc, "CDMX’s collectivos are not content sources for Magma. They are the living network itself. Each should be able to have a distinct home, set its own tone, appoint its own moderators, and retain agency over how it appears.")
add_heading(doc, "A colectivo space includes", 2)
for item in [
    "Identity, members, neighborhood or online context, and its own community rules.",
    "Current projects, concrete needs, open calls, events, radio sessions, and documentation.",
    "A clear way to ask for skills, collaborators, space, equipment, distribution, or care.",
    "Visible shared credit when people or other collectivos make a project possible.",
]:
    add_bullet(doc, item)
add_callout(doc, "Magma’s role", "Surface relevant relationships: a mural colectivo seeking a sound artist, a radio group needing visual identity, or a zine seeking distribution. Magma helps the right edges touch.")

# 5.
doc.add_page_break()
add_heading(doc, "5. Salas and event rooms", 1)
add_lead(doc, "Small rooms make the online house feel inhabited.")
paragraph(doc, "Live chat is valuable when it belongs to a shared moment. Rather than one permanent unrestricted global chat, Magma can host salas connected to radio, collectivos, events, and temporary citywide conversations. They make space for the small, strange, low-stakes interactions that turn strangers into collaborators.")
add_heading(doc, "The event-room rhythm", 2)
add_key_value_table(doc, [
    ("48 hours before", "Introductions, logistics, access information, and the question: who is bringing what?"),
    ("During", "Live notes, photos, music IDs, spontaneous coordination, and shared presence."),
    ("24–72 hours after", "Reflections, links, collaborator follow-ups, and a prompt for what should happen next."),
    ("Afterward", "The room becomes a read-only archive or folds into the event page, preserving a small cultural record."),
])
add_heading(doc, "A focused Art Week pilot", 2)
paragraph(doc, "During Art Week, Casita Magma can be the living community layer around a deliberately small constellation of aligned exhibitions, performances, open studios, and collectivos. Each partner receives a hosted event space, live room, documentation trail, and a clear moderator or host. The archive lets people inside and outside CDMX follow the work without flattening it into a feed.")
add_callout(doc, "Start small", "Invite 5–10 aligned partners, not hundreds. Depth, trust, and a shared point of view are the product.")

# 6.
doc.add_page_break()
add_heading(doc, "6. Community governance", 1)
add_lead(doc, "Open-hearted. Carefully held. Built through participation.")
paragraph(doc, "Magma should not be an imposing cultural authority. It should provide tools and citywide care while the community shapes its own spaces. Moderation needs to be transparent, contextual, and accountable—not absent.")
for item in [
    "Each colectivo appoints or elects moderators for its own space.",
    "Citywide baseline rules protect people from harassment, exploitation, threats, doxxing, and spam.",
    "Reporting, blocking, slow mode, room hosts, and temporary closures protect real-time spaces.",
    "Important moderation decisions are explainable and appealable, with local context respected.",
    "The founder’s role is steward: build tools, protect the values, support moderators, and intervene only for serious cross-community harm.",
]:
    add_bullet(doc, item)

# 7.
doc.add_page_break()
add_heading(doc, "7. From digital commons to cultural house", 1)
add_lead(doc, "The site is not the end product. It is the first room of the house.")
paragraph(doc, "Casita Magma can prove the community a physical home requires before a lease is ever signed: active collectivos, a reliable gathering rhythm, trusted hosts, shared aesthetics, an archive of meaningful work, and people who already help one another.")
add_phase_table(doc)
add_heading(doc, "Immediate next move", 2)
paragraph(doc, "Create one founding manifesto and one first-event brief. Use them to invite a small initial circle of artists, collectivos, and collaborators into a gathering that demonstrates the whole loop: discover, contribute, meet, document, and return to help make the next thing happen.")
add_callout(doc, "North star", "A cultural house built by and for the people making Mexico City.")

doc.save(OUT)
print(OUT)
